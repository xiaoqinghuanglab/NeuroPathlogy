"""extract_variable_matrix.py

Reads per-report, per-seed extracted JSONs from the pipeline output directory,
computes a majority-vote value for each of the 199 NACC NP variables across 5 seeds,
then outputs variable_matrix_{MODEL}_{MODE}.xlsx with three sheets:
  - Variable Summary : per-variable presence rate grouped by schema section
  - Binary Matrix    : reports × variables, 1 = informative, 0 = uninformative/sentinel
  - Raw Values       : reports × variables, raw majority-voted code (None = not extracted)

Tie-breaking: prefer the non-sentinel value. If all tied values are the same type,
pick the first after sorting by string representation.

── MODE SWITCH ──────────────────────────────────────────────────────────────────
Set MODE = "primary"  to process the 199 NACC NP Form variables (.extracted.json).
Set MODE = "residual" to process the 49 residual + 4 programmatic delta variables
           (.residual.json).

Usage: python extract_variable_matrix.py
"""

import json
from collections import Counter
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# ─── CONFIGURE THESE ──────────────────────────────────────────────────────────
MODE         = "residual"          # "primary"  or  "residual"
MODEL        = "oss-20b"
OUTPUT_DIR   = Path("/N/project/ADRD/neuropathoroot/results")
_MODEL_DIR   = Path(f"/N/project/ADRD/neuropathoroot/output/{MODEL}")
INPUT_BASE   = _MODEL_DIR / "residual_extraction" if MODE == "residual" else _MODEL_DIR
N_SEEDS      = 5
# output filename constructed in main() as variable_matrix_{MODEL}_{MODE}.xlsx
# ──────────────────────────────────────────────────────────────────────────────


# ══════════════════════════════════════════════════════════════════════════════
#  PRIMARY MODE — 199 NACC NP Form variables
# ══════════════════════════════════════════════════════════════════════════════

# Per-variable sentinel codes from NACC RDD AllowableCodes.
# No global set — codes like 88 are valid ages for NACCDAGE but sentinel for other vars.
# None (extraction failure) is always treated as uninformative.
SENTINEL_BY_VAR = {
    "NACCAMY":   {8, 9},
    "NACCARTE":  {8, 9},
    "NACCAVAS":  {8, 9},
    "NACCBNKF":  {-4, 9},
    "NACCBRAA":  {8, 9},
    "NACCBRNN":  {8},
    "NACCCBD":   {8, 9},
    "NACCCSFP":  {-4, 9},
    "NACCDAGE":  {888, 999},
    "NACCDIFF":  {8, 9},
    "NACCDOWN":  {7},
    "NACCFORM":  {-4, 9},
    "NACCHEM":   {8, 9},
    "NACCINF":   {8, 9},
    "NACCINT":   {888, 999},
    "NACCLEWY":  {8, 9},
    "NACCMICR":  {8, 9},
    "NACCMOD":   {88, 99},
    "NACCNEC":   {8, 9},
    "NACCNEUR":  {8, 9},
    "NACCOTHP":  {8, 9},
    "NACCPARA":  {-4, 9},
    "NACCPICK":  {8, 9},
    "NACCPRIO":  {8, 9},
    "NACCPROG":  {8, 9},
    "NACCVASC":  {9},
    "NACCYOD":   {8888, 9999},
    "NPABAN":    {-4, 8},
    "NPABANX":   {-4},
    "NPADNC":    {-4, 8, 9},
    "NPADRDA":   {3, 9},
    "NPALSMND":  {-4, 8, 9},
    "NPART":     {-4, 3, 9},
    "NPASAN":    {-4, 8},
    "NPASANX":   {-4},
    "NPBNKB":    {-4, 9},
    "NPBNKF":    {-4, 9},
    "NPCERAD":   {5, 9},
    "NPCHROM":   {50, 99},
    "NPFAUT":    {-4, 9},
    "NPFAUT1":   {-4},
    "NPFAUT2":   {-4},
    "NPFAUT3":   {-4},
    "NPFAUT4":   {-4},
    "NPFIX":     {-4},
    "NPFIXX":    {-4},
    "NPFRONT":   {-4, 3, 9},
    "NPFTD":     {-4, 4, 9},
    "NPFTDNO":   {-4, 3, 9},
    "NPFTDSPC":  {-4, 3, 9},
    "NPFTDT10":  {-4, 8, 9},
    "NPFTDT2":   {-4, 8, 9},
    "NPFTDT5":   {-4, 8, 9},
    "NPFTDT6":   {-4, 8, 9},
    "NPFTDT7":   {-4, 8, 9},
    "NPFTDT8":   {-4, 8, 9},
    "NPFTDT9":   {-4, 8, 9},
    "NPFTDTAU":  {-4, 8, 9},
    "NPFTDTDP":  {-4, 8, 9},
    "NPGENE":    {9},
    "NPGRCCA":   {-4, 8, 9},
    "NPGRHA":    {-4, 8, 9},
    "NPGRLA":    {-4, 8, 9},
    "NPGRLCH":   {-4, 8, 9},
    "NPGRSNH":   {-4, 8, 9},
    "NPHEM":     {-4, 3, 9},
    "NPHEMO":    {-4, 8, 9},
    "NPHEMO1":   {-4, 8, 9},
    "NPHEMO2":   {-4, 8, 9},
    "NPHEMO3":   {-4, 8, 9},
    "NPHIPSCL":  {-4, 8, 9},
    "NPHISG":    {-4},
    "NPHISMB":   {-4},
    "NPHISO":    {-4},
    "NPHISOX":   {-4},
    "NPHISSS":   {-4},
    "NPHIST":    {-4},
    "NPINF":     {-4, 8, 9},
    "NPINF1A":   {-4, 88, 99},
    "NPINF1B":   {-4.4, 88.8, 99.9},
    "NPINF1D":   {-4.4, 88.8, 99.9},
    "NPINF1F":   {-4.4, 88.8, 99.9},
    "NPINF2A":   {-4, 88, 99},
    "NPINF2B":   {-4.4, 88.8, 99.9},
    "NPINF2D":   {-4.4, 88.8, 99.9},
    "NPINF2F":   {-4.4, 88.8, 99.9},
    "NPINF3A":   {-4, 88, 99},
    "NPINF3B":   {-4.4, 88.8, 99.9},
    "NPINF3D":   {-4.4, 88.8, 99.9},
    "NPINF3F":   {-4.4, 88.8, 99.9},
    "NPINF4A":   {-4, 88, 99},
    "NPINF4B":   {-4.4, 88.8, 99.9},
    "NPINF4D":   {-4.4, 88.8, 99.9},
    "NPINF4F":   {-4.4, 88.8, 99.9},
    "NPLAC":     {-4, 3, 9},
    "NPLBOD":    {-4, 8, 9},
    "NPLEWYCS":  {6, 9},
    "NPLINF":    {-4, 3, 9},
    "NPMICRO":   {-4, 3, 9},
    "NPNIT":     {5, 9},
    "NPNLOSS":   {-4, 8, 9},
    "NPOANG":    {-4, 3, 9},
    "NPOCRIT":   {3, 9},
    "NPOFTD":    {-4, 8, 9},
    "NPOFTD1":   {-4, 8, 9},
    "NPOFTD2":   {-4, 8, 9},
    "NPOFTD3":   {-4, 8, 9},
    "NPOFTD4":   {-4, 8, 9},
    "NPOFTD5":   {-4, 8, 9},
    "NPOLD":     {-4, 8, 9},
    "NPOLD1":    {-4, 8, 9},
    "NPOLD2":    {-4, 8, 9},
    "NPOLD3":    {-4, 8, 9},
    "NPOLD4":    {-4, 8, 9},
    "NPOLDD":    {-4, 8, 9},
    "NPOLDD1":   {-4, 8, 9},
    "NPOLDD2":   {-4, 8, 9},
    "NPOLDD3":   {-4, 8, 9},
    "NPOLDD4":   {-4, 8, 9},
    "NPPATH":    {-4, 8, 9},
    "NPPATH10":  {-4, 8, 9},
    "NPPATH11":  {-4, 8, 9},
    "NPPATH2":   {-4, 8, 9},
    "NPPATH3":   {-4, 8, 9},
    "NPPATH4":   {-4, 8, 9},
    "NPPATH5":   {-4, 8, 9},
    "NPPATH6":   {-4, 8, 9},
    "NPPATH7":   {-4, 8, 9},
    "NPPATH8":   {-4, 8, 9},
    "NPPATH9":   {-4, 8, 9},
    "NPPATHO":   {-4},
    "NPPATHOX":  {-4},
    "NPPDXA":    {-4, 8, 9},
    "NPPDXB":    {-4, 8, 9},
    "NPPDXD":    {-4, 8, 9},
    "NPPDXE":    {-4, 8, 9},
    "NPPDXF":    {-4, 8, 9},
    "NPPDXG":    {-4, 8, 9},
    "NPPDXH":    {-4, 8, 9},
    "NPPDXI":    {-4, 8, 9},
    "NPPDXJ":    {-4, 8, 9},
    "NPPDXK":    {-4, 8, 9},
    "NPPDXL":    {-4, 8, 9},
    "NPPDXM":    {-4, 8, 9},
    "NPPDXN":    {-4, 8, 9},
    "NPPDXP":    {-4, 8, 9},
    "NPPDXQ":    {-4, 8, 9},
    "NPPMIH":    {-4, 99.9},
    "NPPRNP":    {9},
    "NPSCL":     {-4, 3, 9},
    "NPTAN":     {-4, 8},
    "NPTANX":    {-4},
    "NPTAU":     {-4, 3, 9},
    "NPTAUHAP":  {9},
    "NPTDPA":    {-4, 8, 9},
    "NPTDPAN":   {-4, 8},
    "NPTDPANX":  {-4},
    "NPTDPB":    {-4, 8, 9},
    "NPTDPC":    {-4, 8, 9},
    "NPTDPD":    {-4, 8, 9},
    "NPTDPE":    {-4, 8, 9},
    "NPTHAL":    {-4, 8, 9},
    "NPVOTH":    {3, 9},
    "NPWBRF":    {-4, 8},
    "NPWBRWT":   {-4, 9999},
    "NPWMR":     {-4, 8, 9},
}

# For character variables: None and empty string are always uninformative.
SENTINEL_CHAR = {None, ""}

# ─── EXCLUSION LIST ───────────────────────────────────────────────────────────
# 1. Admin / logistics — no pathological signal
ADMIN_VARS = {
    "NACCID",       # subject ID (character)
    "NACCADC",      # ADC site number
    "NPFORMVER",    # NP form version
    "NACCMOD",      # month of death
    "NACCYOD",      # year of death
    "NACCINT",      # interval last visit → death
    # tissue banking — logistical, not pathological
    "NACCBNKF",
    "NPBNKB",
    "NACCFORM",
    "NACCPARA",
    "NACCCSFP",
    "NPBNKF",
    # full autopsy flag + free-text organ weights — not brain NP
    "NPFAUT",
}

# 2. Free-text (Character) variables
FREETEXT_VARS = {
    "NPFIXX",       # fixative other specify
    "NPTANX",       # tau antibody other specify
    "NPABANX",      # amyloid beta antibody other specify
    "NPASANX",      # alpha-synuclein antibody other specify
    "NPTDPANX",     # TDP-43 antibody other specify
    "NPHISOX",      # histochemical stain other specify
    "NPPATHOX",     # other vascular pathology specify
    "NACCWRI1",     # write-in diagnosis 1
    "NACCWRI2",     # write-in diagnosis 2
    "NACCWRI3",     # write-in diagnosis 3
    "NPFAUT1",      # full autopsy finding 1
    "NPFAUT2",      # full autopsy finding 2
    "NPFAUT3",      # full autopsy finding 3
    "NPFAUT4",      # full autopsy finding 4
    "NPFHSPEC",     # family history specify
    "NPOTH1X",      # other dx 1 specify
    "NPOTH2X",      # other dx 2 specify
    "NPOTH3X",      # other dx 3 specify
}

# 3. Lab procedure triggers — antibody methods, stain flags; no pathological signal
FREETEXT_TRIGGER_VARS = {
    "NPFIX",        # fixative type — gates NPFIXX; lab procedure, not pathology
    "NPHISO",       # other histochem stain — gates NPHISOX; lab procedure
    "NPWBRF",       # fresh or fixed brain weight — procedural metadata, not pathology
    # Antibody method vars — entirely lab procedure, not pathological signal
    "NPTAN",
    "NPABAN",
    "NPASAN",
    "NPTDPAN",
    # Histochem stain presence flags — lab procedure
    "NPHISMB",
    "NPHISG",
    "NPHISSS",
    "NPHIST",
}

EXCLUDED_VARS = ADMIN_VARS | FREETEXT_VARS | FREETEXT_TRIGGER_VARS

# ─── ORDERED LIST OF 199 VARIABLES ───────────────────────────────────────────
# Must match the order in rddnp.csv. Included vars = these 199 minus EXCLUDED_VARS.
ALL_199_VARS = [
    "NACCID","NACCADC","NPFORMVER","NPSEX","NACCDAGE","NACCMOD","NACCYOD","NACCINT",
    "NPPMIH","NPFIX","NPFIXX","NPWBRWT","NPWBRF","NACCBRNN","NPGRCCA","NPGRLA",
    "NPGRHA","NPGRSNH","NPGRLCH","NACCAVAS","NPTAN","NPTANX","NPABAN","NPABANX",
    "NPASAN","NPASANX","NPTDPAN","NPTDPANX","NPHISMB","NPHISG","NPHISSS","NPHIST",
    "NPHISO","NPHISOX","NPTHAL","NACCBRAA","NACCNEUR","NPADNC","NACCDIFF","NACCVASC",
    "NACCAMY","NPLINF","NPLAC","NPINF","NPINF1A","NPINF1B","NPINF1D","NPINF1F",
    "NPINF2A","NPINF2B","NPINF2D","NPINF2F","NPINF3A","NPINF3B","NPINF3D","NPINF3F",
    "NPINF4A","NPINF4B","NPINF4D","NPINF4F","NACCINF","NPHEM","NPHEMO","NPHEMO1",
    "NPHEMO2","NPHEMO3","NPMICRO","NPOLD","NPOLD1","NPOLD2","NPOLD3","NPOLD4",
    "NACCMICR","NPOLDD","NPOLDD1","NPOLDD2","NPOLDD3","NPOLDD4","NACCHEM","NACCARTE",
    "NPWMR","NPPATH","NACCNEC","NPPATH2","NPPATH3","NPPATH4","NPPATH5","NPPATH6",
    "NPPATH7","NPPATH8","NPPATH9","NPPATH10","NPPATH11","NPPATHO","NPPATHOX","NPART","NPOANG",
    "NACCLEWY","NPLBOD","NPNLOSS","NPHIPSCL","NPSCL","NPFTDTAU","NACCPICK","NPFTDT2",
    "NACCCBD","NACCPROG","NPFTDT5","NPFTDT6","NPFTDT7","NPFTDT8","NPFTDT9","NPFTDT10",
    "NPFRONT","NPTAU","NPFTD","NPFTDTDP","NPALSMND","NPOFTD","NPOFTD1","NPOFTD2",
    "NPOFTD3","NPOFTD4","NPOFTD5","NPFTDNO","NPFTDSPC","NPTDPA","NPTDPB","NPTDPC",
    "NPTDPD","NPTDPE","NPPDXA","NPPDXB","NACCPRIO","NPPDXD","NPPDXE","NPPDXF",
    "NPPDXG","NPPDXH","NPPDXI","NPPDXJ","NPPDXK","NPPDXL","NPPDXM","NPPDXN",
    "NACCDOWN","NPPDXP","NPPDXQ","NACCOTHP","NACCWRI1","NACCWRI2","NACCWRI3",
    "NACCBNKF","NPBNKB","NACCFORM","NACCPARA","NACCCSFP","NPBNKF","NPFAUT",
    "NPFAUT1","NPFAUT2","NPFAUT3","NPFAUT4","NPNIT","NPCERAD","NPADRDA","NPOCRIT",
    "NPVOTH","NPLEWYCS","NPGENE","NPFHSPEC","NPTAUHAP","NPPRNP","NPCHROM",
    "NPPNORM","NPCNORM","NPPADP","NPCADP","NPPAD","NPCAD","NPPLEWY","NPCLEWY",
    "NPPVASC","NPCVASC","NPPFTLD","NPCFTLD","NPPHIPP","NPCHIPP","NPPPRION","NPCPRION",
    "NPPOTH1","NPCOTH1","NPOTH1X","NPPOTH2","NPCOTH2","NPOTH2X","NPPOTH3","NPCOTH3",
    "NPOTH3X",
]

INCLUDED_VARS = [v for v in ALL_199_VARS if v not in EXCLUDED_VARS]


# ══════════════════════════════════════════════════════════════════════════════
#  RESIDUAL MODE — 49 extracted + 4 programmatic delta variables (53 total)
# ══════════════════════════════════════════════════════════════════════════════

# Continuous residual variables — sentinel = 9999.0 (not reported in text).
# Includes all weight blocks, mm measurements, CoW diameters, and deltas.
RESIDUAL_CONTINUOUS_VARS = {
    # bilateral_hemibrain_weights
    "hemibrain_weight_right_fresh_g",
    "hemibrain_weight_left_fresh_g",
    "hemibrain_weight_right_fixed_g",
    "hemibrain_weight_left_fixed_g",
    # bilateral_cerebral_weights
    "cerebral_weight_right_fresh_g",
    "cerebral_weight_left_fresh_g",
    "cerebral_weight_right_fixed_g",
    "cerebral_weight_left_fixed_g",
    # bilateral_cerebellar_weights
    "cerebellar_weight_right_fresh_g",
    "cerebellar_weight_left_fresh_g",
    "cerebellar_weight_right_fixed_g",
    "cerebellar_weight_left_fixed_g",
    # bilateral_brainstem_weights
    "brainstem_weight_right_fresh_g",
    "brainstem_weight_left_fresh_g",
    "brainstem_weight_right_fixed_g",
    "brainstem_weight_left_fixed_g",
    # corpus_callosum_measurements
    "corpus_callosum_genu_mm",
    "corpus_callosum_body_anterior_mm",
    "corpus_callosum_body_mid_mm",
    "corpus_callosum_body_posterior_mm",
    "corpus_callosum_splenium_mm",
    # circle_of_willis_diameters
    "cow_basilar_mm",
    "cow_vertebral_right_mm",
    "cow_vertebral_left_mm",
    "cow_ica_right_mm",
    "cow_ica_left_mm",
    "cow_mca_right_mm",
    "cow_mca_left_mm",
    "cow_aca_right_mm",
    "cow_aca_left_mm",
    "cow_pca_right_mm",
    "cow_pca_left_mm",
    "cow_pcom_right_mm",
    "cow_pcom_left_mm",
    "cow_acom_mm",
    # structural_measurements
    "caudate_nucleus_head_width_mm",
    # programmatic deltas — computed post-extraction, not LLM-extracted
    "hemibrain_weight_delta_g",
    "cerebral_weight_delta_g",
    "cerebellar_weight_delta_g",
    "brainstem_weight_delta_g",
}

# Ordinal residual variables — sentinel = 9 (not mentioned in report).
# Code 0 = examined and absent/normal (informative). Code 8 = not assessed.
RESIDUAL_ORDINAL_VARS = {
    # structural_severity_markers
    "lateral_ventricle_enlargement_severity",
    "amygdala_atrophy_severity",
    "dilated_perivascular_spaces",
    "cerebellar_atrophy_severity",
    "globus_pallidus_neuronal_loss_severity",
    "basal_ganglia_atrophy",
    "thalamic_degeneration_severity",
    "brainstem_atrophy_severity",
    # regional_pathology_severity
    "purkinje_cell_loss_severity",
    "dentate_nucleus_atrophy",
    "artag_severity",
    "gvd_hippocampus_severity",
    "hirano_bodies_hippocampus_severity",
}

# ─── ORDERED LIST OF 53 RESIDUAL VARIABLES ───────────────────────────────────
# 49 LLM-extracted + 4 programmatic deltas appended at the end.
# Must match the block order in residual_extract.py / ResidualOutput schema.
ALL_RESIDUAL_VARS = [
    # bilateral_hemibrain_weights (4)
    "hemibrain_weight_right_fresh_g",
    "hemibrain_weight_left_fresh_g",
    "hemibrain_weight_right_fixed_g",
    "hemibrain_weight_left_fixed_g",
    # bilateral_cerebral_weights (4)
    "cerebral_weight_right_fresh_g",
    "cerebral_weight_left_fresh_g",
    "cerebral_weight_right_fixed_g",
    "cerebral_weight_left_fixed_g",
    # bilateral_cerebellar_weights (4)
    "cerebellar_weight_right_fresh_g",
    "cerebellar_weight_left_fresh_g",
    "cerebellar_weight_right_fixed_g",
    "cerebellar_weight_left_fixed_g",
    # bilateral_brainstem_weights (4)
    "brainstem_weight_right_fresh_g",
    "brainstem_weight_left_fresh_g",
    "brainstem_weight_right_fixed_g",
    "brainstem_weight_left_fixed_g",
    # corpus_callosum_measurements (5)
    "corpus_callosum_genu_mm",
    "corpus_callosum_body_anterior_mm",
    "corpus_callosum_body_mid_mm",
    "corpus_callosum_body_posterior_mm",
    "corpus_callosum_splenium_mm",
    # circle_of_willis_diameters (14)
    "cow_basilar_mm",
    "cow_vertebral_right_mm",
    "cow_vertebral_left_mm",
    "cow_ica_right_mm",
    "cow_ica_left_mm",
    "cow_mca_right_mm",
    "cow_mca_left_mm",
    "cow_aca_right_mm",
    "cow_aca_left_mm",
    "cow_pca_right_mm",
    "cow_pca_left_mm",
    "cow_pcom_right_mm",
    "cow_pcom_left_mm",
    "cow_acom_mm",
    # structural_measurements (1)
    "caudate_nucleus_head_width_mm",
    # structural_severity_markers (8)
    "lateral_ventricle_enlargement_severity",
    "amygdala_atrophy_severity",
    "dilated_perivascular_spaces",
    "cerebellar_atrophy_severity",
    "globus_pallidus_neuronal_loss_severity",
    "basal_ganglia_atrophy",
    "thalamic_degeneration_severity",
    "brainstem_atrophy_severity",
    # regional_pathology_severity (5)
    "purkinje_cell_loss_severity",
    "dentate_nucleus_atrophy",
    "artag_severity",
    "gvd_hippocampus_severity",
    "hirano_bodies_hippocampus_severity",
    # programmatic deltas (4) — appended after all LLM-extracted variables
    "hemibrain_weight_delta_g",
    "cerebral_weight_delta_g",
    "cerebellar_weight_delta_g",
    "brainstem_weight_delta_g",
]

# ─── SCHEMA SECTION GROUPING — RESIDUAL ──────────────────────────────────────
# Mirrors the block structure in residual_extract.py.
# Deltas grouped separately so it's clear they are computed, not extracted.
#
# Section ordering:
#   1  Hemibrain Weights   — bilateral_hemibrain_weights block
#   2  Cerebral Weights    — bilateral_cerebral_weights block
#   3  Cerebellar Weights  — bilateral_cerebellar_weights block
#   4  Brainstem Weights   — bilateral_brainstem_weights block
#   5  Weight Deltas       — programmatic asymmetry deltas (right − left)
#   6  Corpus Callosum     — corpus_callosum_measurements block
#   7  Circle of Willis    — circle_of_willis_diameters block
#   8  Structural Measures — structural_measurements block
#   9  Structural Severity — structural_severity_markers block
#  10  Regional Pathology  — regional_pathology_severity block

RESIDUAL_SECTION_ORDER = {
    "Hemibrain Weights":    1,
    "Cerebral Weights":     2,
    "Cerebellar Weights":   3,
    "Brainstem Weights":    4,
    "Weight Deltas":        5,
    "Corpus Callosum":      6,
    "Circle of Willis":     7,
    "Structural Measures":  8,
    "Structural Severity":  9,
    "Regional Pathology":  10,
}

RESIDUAL_SECTION_MAP = {
    # ── Hemibrain Weights (bilateral_hemibrain_weights) ──────────────────────
    "hemibrain_weight_right_fresh_g":  "Hemibrain Weights",
    "hemibrain_weight_left_fresh_g":   "Hemibrain Weights",
    "hemibrain_weight_right_fixed_g":  "Hemibrain Weights",
    "hemibrain_weight_left_fixed_g":   "Hemibrain Weights",
    # ── Cerebral Weights (bilateral_cerebral_weights) ────────────────────────
    "cerebral_weight_right_fresh_g":   "Cerebral Weights",
    "cerebral_weight_left_fresh_g":    "Cerebral Weights",
    "cerebral_weight_right_fixed_g":   "Cerebral Weights",
    "cerebral_weight_left_fixed_g":    "Cerebral Weights",
    # ── Cerebellar Weights (bilateral_cerebellar_weights) ────────────────────
    "cerebellar_weight_right_fresh_g": "Cerebellar Weights",
    "cerebellar_weight_left_fresh_g":  "Cerebellar Weights",
    "cerebellar_weight_right_fixed_g": "Cerebellar Weights",
    "cerebellar_weight_left_fixed_g":  "Cerebellar Weights",
    # ── Brainstem Weights (bilateral_brainstem_weights) ──────────────────────
    "brainstem_weight_right_fresh_g":  "Brainstem Weights",
    "brainstem_weight_left_fresh_g":   "Brainstem Weights",
    "brainstem_weight_right_fixed_g":  "Brainstem Weights",
    "brainstem_weight_left_fixed_g":   "Brainstem Weights",
    # ── Weight Deltas (programmatic — right minus left) ──────────────────────
    "hemibrain_weight_delta_g":        "Weight Deltas",
    "cerebral_weight_delta_g":         "Weight Deltas",
    "cerebellar_weight_delta_g":       "Weight Deltas",
    "brainstem_weight_delta_g":        "Weight Deltas",
    # ── Corpus Callosum (corpus_callosum_measurements) ───────────────────────
    "corpus_callosum_genu_mm":              "Corpus Callosum",
    "corpus_callosum_body_anterior_mm":     "Corpus Callosum",
    "corpus_callosum_body_mid_mm":          "Corpus Callosum",
    "corpus_callosum_body_posterior_mm":    "Corpus Callosum",
    "corpus_callosum_splenium_mm":          "Corpus Callosum",
    # ── Circle of Willis (circle_of_willis_diameters) ────────────────────────
    "cow_basilar_mm":        "Circle of Willis",
    "cow_vertebral_right_mm":"Circle of Willis",
    "cow_vertebral_left_mm": "Circle of Willis",
    "cow_ica_right_mm":      "Circle of Willis",
    "cow_ica_left_mm":       "Circle of Willis",
    "cow_mca_right_mm":      "Circle of Willis",
    "cow_mca_left_mm":       "Circle of Willis",
    "cow_aca_right_mm":      "Circle of Willis",
    "cow_aca_left_mm":       "Circle of Willis",
    "cow_pca_right_mm":      "Circle of Willis",
    "cow_pca_left_mm":       "Circle of Willis",
    "cow_pcom_right_mm":     "Circle of Willis",
    "cow_pcom_left_mm":      "Circle of Willis",
    "cow_acom_mm":           "Circle of Willis",
    # ── Structural Measures (structural_measurements) ────────────────────────
    "caudate_nucleus_head_width_mm":   "Structural Measures",
    # ── Structural Severity (structural_severity_markers) ────────────────────
    "lateral_ventricle_enlargement_severity": "Structural Severity",
    "amygdala_atrophy_severity":       "Structural Severity",
    "dilated_perivascular_spaces":     "Structural Severity",
    "cerebellar_atrophy_severity":     "Structural Severity",
    "globus_pallidus_neuronal_loss_severity": "Structural Severity",
    "basal_ganglia_atrophy":           "Structural Severity",
    "thalamic_degeneration_severity":  "Structural Severity",
    "brainstem_atrophy_severity":      "Structural Severity",
    # ── Regional Pathology (regional_pathology_severity) ─────────────────────
    "purkinje_cell_loss_severity":     "Regional Pathology",
    "dentate_nucleus_atrophy":         "Regional Pathology",
    "artag_severity":                  "Regional Pathology",
    "gvd_hippocampus_severity":        "Regional Pathology",
    "hirano_bodies_hippocampus_severity": "Regional Pathology",
}


# ══════════════════════════════════════════════════════════════════════════════
#  SHARED UTILITIES
# ══════════════════════════════════════════════════════════════════════════════

def flatten_json(d, out=None):
    """Recursively flatten nested JSON dict into {variable: value}, skipping
    field_annotations, extraction_confidence, extraction_notes blocks."""
    if out is None:
        out = {}
    skip_keys = {"field_annotations", "extraction_confidence", "extraction_notes"}
    for k, v in d.items():
        if k in skip_keys:
            continue
        if isinstance(v, dict):
            flatten_json(v, out)
        else:
            out[k] = v
    return out


def flatten_residual_json(d):
    """Flatten .residual.json into a single {variable: value} dict.

    Pulls variables from the 9 named blocks and the top-level 'programmatic'
    key separately. Skips metadata keys. One level of nesting only — matches
    the output structure of residual_extract.py / ResidualOutput.
    """
    RESIDUAL_BLOCKS = {
        "bilateral_hemibrain_weights", "bilateral_cerebral_weights",
        "bilateral_cerebellar_weights", "bilateral_brainstem_weights",
        "corpus_callosum_measurements", "circle_of_willis_diameters",
        "structural_measurements", "structural_severity_markers",
        "regional_pathology_severity",
    }
    SKIP_KEYS = {"field_annotations", "extraction_confidence", "extraction_notes"}
    out = {}
    for key, val in d.items():
        if key in SKIP_KEYS:
            continue
        if key in RESIDUAL_BLOCKS and isinstance(val, dict):
            out.update(val)
        elif key == "programmatic" and isinstance(val, dict):
            out.update(val)
        # top-level scalars outside the above (unexpected) are ignored
    return out


def is_sentinel(value, var_name=None):
    """Return True if the value is uninformative (primary mode).
    Uses SENTINEL_BY_VAR derived directly from the RDD — no global sentinel set."""
    if value is None:
        return True
    if isinstance(value, str):
        s = value.strip()
        return s in SENTINEL_CHAR or s == ""
    try:
        num = float(value)
        sentinels = SENTINEL_BY_VAR.get(var_name, set()) if var_name else set()
        return round(num, 1) in sentinels
    except (ValueError, TypeError):
        return False


def is_sentinel_residual(value, var_name=None):
    """Return True if the value is uninformative (residual mode).

    Continuous variables (weights, mm, deltas): sentinel = 9999.0.
    Ordinal variables (severity scales 0-3/8/9): sentinel = 9 (not mentioned).
      Note: code 0 (examined, absent/normal) IS informative — do not treat as sentinel.
    None (extraction failure) is always uninformative.
    """
    if value is None:
        return True
    try:
        num = float(value)
    except (ValueError, TypeError):
        return True
    if var_name in RESIDUAL_CONTINUOUS_VARS:
        return round(num, 1) == 9999.0
    if var_name in RESIDUAL_ORDINAL_VARS:
        return round(num) == 9
    # fallback for any variable not explicitly categorised
    return round(num, 1) == 9999.0


def majority_vote(values, var_name=None, sentinel_fn=None):
    """
    Given a list of 5 values (may include None), return the majority value.
    Tie-breaking: prefer the non-sentinel (informative) value.
    If all tied values are the same type (all sentinel or all informative),
    return the first after sorting by string representation.

    sentinel_fn: callable(value, var_name) -> bool — passed explicitly by
    compute_row / compute_raw_row so this function never reads MODE directly.
    """
    # Filter out Python None from vote (None = extraction failure, not a real code)
    non_none = [v for v in values if v is not None]
    if not non_none:
        return None

    counts = Counter(str(v) for v in non_none)
    max_count = max(counts.values())
    candidates_str = [k for k, c in counts.items() if c == max_count]

    if len(candidates_str) == 1:
        # Unambiguous majority — recover original typed value
        winner_str = candidates_str[0]
        for v in non_none:
            if str(v) == winner_str:
                return v
        return winner_str

    # Tie — prefer non-sentinel
    str_to_val = {}
    for v in non_none:
        str_to_val[str(v)] = v

    candidate_vals = [str_to_val[s] for s in candidates_str]
    if sentinel_fn is not None:
        informative = [v for v in candidate_vals if not sentinel_fn(v, var_name)]
        if informative:
            return sorted(informative, key=lambda x: str(x))[0]

    # All candidates are sentinel (or no sentinel_fn provided) — pick first alphabetically
    return sorted(candidate_vals, key=lambda x: str(x))[0]


def get_report_id(filename):
    """Extract report ID from filename.

    Primary:  '1994-630_3129_seed0.extracted.json' → '1994-630_3129'
    Residual: '6472_seed0.residual.json'           → '6472'
    """
    stem = Path(filename).stem  # strips .json
    # strip the inner extension (.extracted or .residual)
    for ext in (".extracted", ".residual"):
        if stem.endswith(ext):
            stem = stem[: -len(ext)]
            break
    # strip _seedN suffix
    parts = stem.rsplit("_seed", 1)
    return parts[0]


def load_seed_files(input_dir):
    """
    Scan input_dir for seed files matching the active MODE.
    Returns dict: {report_id: [flat_dict_seed0, flat_dict_seed1, ...]}
    """
    input_dir = Path(input_dir)
    reports = {}

    if MODE == "primary":
        pattern    = "*.extracted.json"
        flatten_fn = flatten_json
    else:
        pattern    = "*.residual.json"
        flatten_fn = flatten_residual_json

    for fpath in sorted(input_dir.glob(pattern)):
        report_id = get_report_id(fpath.name)
        flat = flatten_fn(json.loads(fpath.read_text(encoding="utf-8")))
        reports.setdefault(report_id, []).append(flat)
    return reports


def _active_vars():
    """Return the included variable list for the active MODE."""
    return INCLUDED_VARS if MODE == "primary" else ALL_RESIDUAL_VARS


def _active_sentinel_fn():
    """Return the sentinel check function for the active MODE."""
    return is_sentinel if MODE == "primary" else is_sentinel_residual


def compute_row(report_id, seed_dicts):
    """
    For one report, majority-vote each included variable across seeds.
    Returns dict: {variable: 1 or 0}
    """
    sentinel_fn = _active_sentinel_fn()
    row = {"report_id": report_id}
    for var in _active_vars():
        values = [d.get(var) for d in seed_dicts]
        majority = majority_vote(values, var_name=var, sentinel_fn=sentinel_fn)
        row[var] = 0 if sentinel_fn(majority, var_name=var) else 1
    return row


def compute_raw_row(report_id, seed_dicts):
    """
    For one report, majority-vote each included variable across seeds.
    Returns dict: {variable: raw majority value (int/float/None as-is)}
    """
    sentinel_fn = _active_sentinel_fn()
    row = {"report_id": report_id}
    for var in _active_vars():
        values = [d.get(var) for d in seed_dicts]
        majority = majority_vote(values, var_name=var, sentinel_fn=sentinel_fn)
        # Store the raw value — None if all seeds returned None
        # Numeric values are kept as int where possible for clean Excel output
        if majority is None:
            row[var] = None
        elif isinstance(majority, float) and majority == int(majority):
            row[var] = int(majority)
        else:
            row[var] = majority
    return row


# ─── SCHEMA SECTION GROUPING — PRIMARY ───────────────────────────────────────
# Verified against the JSON output structure (see sample JSON).
# Ordered to read logically: case info → pathology → vascular → diagnoses → misc.
#
# Section ordering (for sort key):
#   1  Specimen Info       — case_metadata + specimen_info blocks
#   2  Gross Exam          — gross_and_vascular block (gross findings only)
#   3  AD Pathology        — ad_pathology block
#   4  Lewy Body           — lewy_pathology block
#   5  Microscopic         — microscopic_findings block (neuronal loss, hippo sclerosis,
#                            medial temporal sclerosis — NPSCL moved here from Other)
#   6  Infarcts            — infarct_parent + infarct_counts + infarct_sizes blocks
#   7  Hemorrhage          — hemorrhage block
#   8  Microinfarcts       — microinfarcts block
#   9  Microbleeds         — microbleeds block
#  10  Other Vascular      — misc_vascular + other_vascular blocks
#                            NACCVASC moved here (computed vascular summary)
#  11  FTLD Tau            — ftld_tau block
#  12  FTLD General        — ftld block + other_ftld + tdp_flag + prion_flag
#  13  TDP-43              — tdp43_distribution block
#  14  Diagnoses           — primary_contrib_dx + normal_and_insuff_ad +
#                            hippocampal_and_prion_dx + other_dx_pairs + write_in_dx
#                            NACCPRIO moved here (prion disease flag)
#  15  Genetics            — genetics + family_history blocks
#  16  Other Disease Flags — other_disease_flags block (NPPDX*, NACCDOWN)
#  17  Criteria & Misc     — criteria_and_misc block

SECTION_ORDER = {
    "Specimen Info":       1,
    "Gross Exam":          2,
    "AD Pathology":        3,
    "Lewy Body":           4,
    "Microscopic":         5,
    "Infarcts":            6,
    "Hemorrhage":          7,
    "Microinfarcts":       8,
    "Microbleeds":         9,
    "Other Vascular":     10,
    "FTLD Tau":           11,
    "FTLD General":       12,
    "TDP-43":             13,
    "Diagnoses":          14,
    "Genetics":           15,
    "Other Disease Flags":16,
    "Criteria & Misc":    17,
}

SECTION_MAP = {
    # ── Specimen Info (specimen_info + case_metadata) ────────────────────────
    "NPSEX":    "Specimen Info",
    "NACCDAGE": "Specimen Info",
    "NPPMIH":   "Specimen Info",
    "NPWBRWT":  "Specimen Info",
    "NPWBRF":   "Specimen Info",
    # ── Gross Exam (gross_and_vascular — gross findings) ────────────────────
    "NPGRCCA":  "Gross Exam",
    "NPGRLA":   "Gross Exam",
    "NPGRHA":   "Gross Exam",
    "NPGRSNH":  "Gross Exam",
    "NPGRLCH":  "Gross Exam",
    "NACCAVAS": "Gross Exam",
    "NPWMR":    "Gross Exam",
    "NACCARTE": "Gross Exam",
    # ── AD Pathology (ad_pathology) ─────────────────────────────────────────
    "NPTHAL":   "AD Pathology",
    "NACCBRAA": "AD Pathology",
    "NACCNEUR": "AD Pathology",
    "NPADNC":   "AD Pathology",
    "NACCDIFF": "AD Pathology",
    "NACCAMY":  "AD Pathology",
    # ── Lewy Body (lewy_pathology) ──────────────────────────────────────────
    "NPLBOD":   "Lewy Body",
    "NACCLEWY": "Lewy Body",
    # ── Microscopic (microscopic_findings) ──────────────────────────────────
    # NPSCL = medial temporal lobe sclerosis — belongs here, not Other Disease Flags
    "NACCBRNN": "Microscopic",
    "NPNLOSS":  "Microscopic",
    "NPHIPSCL": "Microscopic",
    "NPSCL":    "Microscopic",
    # ── Infarcts (infarct_parent + infarct_counts + infarct_sizes) ──────────
    "NPLINF":   "Infarcts",
    "NPLAC":    "Infarcts",
    "NPINF":    "Infarcts",
    "NACCINF":  "Infarcts",
    "NPINF1A":  "Infarcts", "NPINF1B": "Infarcts", "NPINF1D": "Infarcts", "NPINF1F": "Infarcts",
    "NPINF2A":  "Infarcts", "NPINF2B": "Infarcts", "NPINF2D": "Infarcts", "NPINF2F": "Infarcts",
    "NPINF3A":  "Infarcts", "NPINF3B": "Infarcts", "NPINF3D": "Infarcts", "NPINF3F": "Infarcts",
    "NPINF4A":  "Infarcts", "NPINF4B": "Infarcts", "NPINF4D": "Infarcts", "NPINF4F": "Infarcts",
    # ── Hemorrhage (hemorrhage) ─────────────────────────────────────────────
    "NPHEM":    "Hemorrhage",
    "NPHEMO":   "Hemorrhage",
    "NPHEMO1":  "Hemorrhage",
    "NPHEMO2":  "Hemorrhage",
    "NPHEMO3":  "Hemorrhage",
    "NACCHEM":  "Hemorrhage",
    # ── Microinfarcts (microinfarcts) ────────────────────────────────────────
    "NPOLD":    "Microinfarcts",
    "NPOLD1":   "Microinfarcts",
    "NPOLD2":   "Microinfarcts",
    "NPOLD3":   "Microinfarcts",
    "NPOLD4":   "Microinfarcts",
    "NACCMICR": "Microinfarcts",
    # ── Microbleeds (microbleeds) ────────────────────────────────────────────
    "NPOLDD":   "Microbleeds",
    "NPOLDD1":  "Microbleeds",
    "NPOLDD2":  "Microbleeds",
    "NPOLDD3":  "Microbleeds",
    "NPOLDD4":  "Microbleeds",
    # ── Other Vascular (misc_vascular + other_vascular + NACCVASC) ──────────
    # misc_vascular: NPMICRO, NPART, NPOANG
    # other_vascular: NPPATH, NPPATH2-11, NACCNEC, NPPATHO
    # NACCVASC: computed vascular summary — placed here with its source variables
    "NPMICRO":  "Other Vascular",
    "NPART":    "Other Vascular",
    "NPOANG":   "Other Vascular",
    "NPPATH":   "Other Vascular",
    "NACCNEC":  "Other Vascular",
    "NPPATH2":  "Other Vascular", "NPPATH3":  "Other Vascular", "NPPATH4":  "Other Vascular",
    "NPPATH5":  "Other Vascular", "NPPATH6":  "Other Vascular", "NPPATH7":  "Other Vascular",
    "NPPATH8":  "Other Vascular", "NPPATH9":  "Other Vascular", "NPPATH10": "Other Vascular",
    "NPPATH11": "Other Vascular",
    "NPPATHO":  "Other Vascular",
    "NACCVASC": "Other Vascular",
    # ── FTLD Tau (ftld_tau) ─────────────────────────────────────────────────
    "NPFTDTAU": "FTLD Tau",
    "NACCPICK": "FTLD Tau",
    "NPFTDT2":  "FTLD Tau",
    "NACCCBD":  "FTLD Tau",
    "NACCPROG": "FTLD Tau",
    "NPFTDT5":  "FTLD Tau",
    "NPFTDT6":  "FTLD Tau",
    "NPFTDT7":  "FTLD Tau",
    "NPFTDT8":  "FTLD Tau",
    "NPFTDT9":  "FTLD Tau",
    "NPFTDT10": "FTLD Tau",
    # ── FTLD General (ftld + other_ftld + tdp_flag + prion_flag) ────────────
    "NPFTD":    "FTLD General",
    "NPTAU":    "FTLD General",
    "NPFRONT":  "FTLD General",
    "NPFTDTDP": "FTLD General",
    "NPALSMND": "FTLD General",
    "NPOFTD":   "FTLD General",
    "NPOFTD1":  "FTLD General",
    "NPOFTD2":  "FTLD General",
    "NPOFTD3":  "FTLD General",
    "NPOFTD4":  "FTLD General",
    "NPOFTD5":  "FTLD General",
    "NPFTDNO":  "FTLD General",
    "NPFTDSPC": "FTLD General",
    # ── TDP-43 (tdp43_distribution) ─────────────────────────────────────────
    "NPTDPA":   "TDP-43",
    "NPTDPB":   "TDP-43",
    "NPTDPC":   "TDP-43",
    "NPTDPD":   "TDP-43",
    "NPTDPE":   "TDP-43",
    # ── Diagnoses (primary_contrib_dx + normal_and_insuff_ad + ──────────────
    #              hippocampal_and_prion_dx + other_dx_pairs + write_in_dx)
    # NACCPRIO = prion disease flag — moved here from Other Disease Flags
    "NPPAD":    "Diagnoses", "NPCAD":    "Diagnoses",
    "NPPLEWY":  "Diagnoses", "NPCLEWY":  "Diagnoses",
    "NPPVASC":  "Diagnoses", "NPCVASC":  "Diagnoses",
    "NPPFTLD":  "Diagnoses", "NPCFTLD":  "Diagnoses",
    "NPPNORM":  "Diagnoses", "NPCNORM":  "Diagnoses",
    "NPPADP":   "Diagnoses", "NPCADP":   "Diagnoses",
    "NPPHIPP":  "Diagnoses", "NPCHIPP":  "Diagnoses",
    "NPPPRION": "Diagnoses", "NPCPRION": "Diagnoses",
    "NACCPRIO": "Diagnoses",
    "NPPOTH1":  "Diagnoses", "NPCOTH1":  "Diagnoses",
    "NPPOTH2":  "Diagnoses", "NPCOTH2":  "Diagnoses",
    "NPPOTH3":  "Diagnoses", "NPCOTH3":  "Diagnoses",
    "NACCOTHP": "Diagnoses",
    # ── Genetics (genetics + family_history) ─────────────────────────────────
    "NPGENE":   "Genetics",
    "NPTAUHAP": "Genetics",
    "NPPRNP":   "Genetics",
    "NPCHROM":  "Genetics",
    "NPPDXP":   "Genetics",   # chromosomal abnormality present
    "NPPDXQ":   "Genetics",   # FTLD-associated gene mutation present
    # ── Other Disease Flags (other_disease_flags: NPPDX* excluding P/Q) ─────
    "NPPDXA":   "Other Disease Flags",
    "NPPDXB":   "Other Disease Flags",
    "NPPDXD":   "Other Disease Flags",
    "NPPDXE":   "Other Disease Flags",
    "NPPDXF":   "Other Disease Flags",
    "NPPDXG":   "Other Disease Flags",
    "NPPDXH":   "Other Disease Flags",
    "NPPDXI":   "Other Disease Flags",
    "NPPDXJ":   "Other Disease Flags",
    "NPPDXK":   "Other Disease Flags",
    "NPPDXL":   "Other Disease Flags",
    "NPPDXM":   "Other Disease Flags",
    "NPPDXN":   "Other Disease Flags",
    "NACCDOWN": "Other Disease Flags",
    # ── Criteria & Misc (criteria_and_misc) ─────────────────────────────────
    "NPNIT":    "Criteria & Misc",
    "NPCERAD":  "Criteria & Misc",
    "NPADRDA":  "Criteria & Misc",
    "NPOCRIT":  "Criteria & Misc",
    "NPVOTH":   "Criteria & Misc",
    "NPLEWYCS": "Criteria & Misc",
}

def get_section(var):
    """Return the section name for a variable in the active MODE."""
    if MODE == "primary":
        return SECTION_MAP.get(var, "Other")
    return RESIDUAL_SECTION_MAP.get(var, "Other")


def build_variable_summary(rows, n_reports):
    """
    For each included variable compute:
      section, n_present, n_absent, pct_present, category.

    Categories (4 tiers — thresholds chosen for 161-report dataset):
      Always Present  : 100% of reports
      Present         : >15% to <100%
      Rarely Present  : >0% to ≤15%
      Never Present   : 0% of reports
    """
    summary = []
    for var in _active_vars():
        counts = [row[var] for row in rows]
        n_present = sum(counts)
        pct = round(n_present / n_reports * 100, 1)

        if pct == 100.0:
            category = "Always Present"
        elif pct == 0.0:
            category = "Never Present"
        elif pct <= 15.0:
            category = "Rarely Present"
        else:
            category = "Present"

        summary.append({
            "Variable":  var,
            "Section":   get_section(var),
            "N Present": n_present,
            "N Absent":  n_reports - n_present,
            "% Present": pct,
            "Category":  category,
        })

    # Sort by section order then % present descending within section
    active_order = SECTION_ORDER if MODE == "primary" else RESIDUAL_SECTION_ORDER
    summary.sort(key=lambda x: (
        active_order.get(x["Section"], 99),
        -x["% Present"]
    ))
    return summary


def write_excel(rows, raw_rows, summary, n_reports, output_path):
    wb = Workbook()

    # ── Colours — matching compare_models_gt.py palette exactly ─────────────
    HDR_DARK    = "1C2833"    # near-black header (same as compare_models_gt)
    COL_PRESENT = "D5F5E3"    # light green cell  (GT match green from compare script)
    COL_ABSENT  = "FFFFFF"    # white
    COL_TOTAL   = "D6E4F0"    # OSS model light blue (MODEL_FILL_HEX["oss"])
    COL_SUMROW  = "EBF5FB"    # summary row tint (same as compare script section 5)
    COL_SECHDR  = "2C3E50"    # section divider (same as compare script _sec_hdr)

    # Category colours — 4 tiers, matching compare_models_gt.py palette
    CAT_FILL = {
        "Always Present": "D5F5E3",   # GT match green
        "Present":        "EBF5FB",   # section-5 blue tint
        "Rarely Present": "FEF3E2",   # TIE_FILL amber
        "Never Present":  "FADBD8",   # GT mismatch red
    }
    CAT_FONT = {
        "Always Present": "145A32",   # dark green
        "Present":        "1A5276",   # dark blue
        "Rarely Present": "784212",   # dark amber
        "Never Present":  "922B21",   # dark red
    }

    thin   = Side(style="thin",   color="CCCCCC")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    def hdr_font(size=10):
        return Font(name="Arial", bold=True, color="FFFFFF", size=size)

    def cell_font(bold=False, size=9, color="000000"):
        return Font(name="Arial", bold=bold, size=size, color=color)

    def fill(hex_color):
        return PatternFill("solid", fgColor=hex_color)

    def ctr():
        return Alignment(horizontal="center", vertical="center", wrap_text=False)

    def ctr_wrap():
        return Alignment(horizontal="center", vertical="center", wrap_text=True)

    def lft():
        return Alignment(horizontal="left", vertical="center")

    # ════════════════════════════════════════════════════════════════════════
    # SHEET 1 — Variable Summary
    ws_sum = wb.active
    ws_sum.title = "Variable Summary"
    ws_sum.sheet_view.showGridLines = False

    sum_headers = ["Variable", "N Present", "N Absent", "% Present", "Category"]
    ws_sum.row_dimensions[1].height = 28
    for c, h in enumerate(sum_headers, 1):
        cell = ws_sum.cell(row=1, column=c, value=h)
        cell.fill = fill(HDR_DARK)
        cell.font = hdr_font(size=9)
        cell.alignment = ctr_wrap()
        cell.border = border

    current_section = None
    r = 2
    for entry in summary:
        # Section divider row — spans all 5 columns, dark background, white bold text
        if entry["Section"] != current_section:
            current_section = entry["Section"]
            ws_sum.row_dimensions[r].height = 16
            for c in range(1, len(sum_headers) + 1):
                cell = ws_sum.cell(row=r, column=c,
                                value=current_section if c == 1 else "")
                cell.fill = fill(COL_SECHDR)
                cell.font = Font(name="Arial", bold=True, size=9, color="FFFFFF")
                cell.alignment = lft() if c == 1 else ctr()
                cell.border = border
            r += 1

        cat = entry["Category"]
        ws_sum.row_dimensions[r].height = 15
        for c, field in enumerate(sum_headers, 1):
            # "Section" key is still in entry dict but not a column anymore
            val = entry[field]
            cell = ws_sum.cell(row=r, column=c, value=val)
            cell.fill = fill(CAT_FILL[cat])
            cell.font = Font(name="Arial", size=9, color=CAT_FONT[cat])
            cell.alignment = lft() if field == "Variable" else ctr()
            cell.border = border
        r += 1

    # Residual variable names are longer — widen column A accordingly
    ws_sum.column_dimensions["A"].width = 38 if MODE == "residual" else 16
    ws_sum.column_dimensions["B"].width = 12
    ws_sum.column_dimensions["C"].width = 12
    ws_sum.column_dimensions["D"].width = 12
    ws_sum.column_dimensions["E"].width = 18
    ws_sum.freeze_panes = "A2"

    # ════════════════════════════════════════════════════════════════════════
    # SHEET 2 — Binary Matrix
    # ════════════════════════════════════════════════════════════════════════
    ws_bin = wb.create_sheet("Binary Matrix")
    ws_bin.sheet_view.showGridLines = False

    active_vars = _active_vars()
    col_fields = ["report_id"] + active_vars + ["total_informative"]

    # Header row — horizontal, single line, fixed narrow columns
    ws_bin.row_dimensions[1].height = 15
    for c, h in enumerate(col_fields, 1):
        cell = ws_bin.cell(row=1, column=c, value=h)
        cell.fill = fill(HDR_DARK)
        cell.font = hdr_font(size=8)
        cell.alignment = Alignment(horizontal="center", vertical="center",
                                   wrap_text=False)
        cell.border = border
        if h == "report_id":
            ws_bin.column_dimensions[get_column_letter(c)].width = 22
        elif h == "total_informative":
            ws_bin.column_dimensions[get_column_letter(c)].width = 12
        else:
            ws_bin.column_dimensions[get_column_letter(c)].width = 9

    # Data rows
    last_data_row = 1
    for r, row in enumerate(rows, 2):
        last_data_row = r
        ws_bin.row_dimensions[r].height = 15
        total = sum(row[v] for v in active_vars)
        for c, field in enumerate(col_fields, 1):
            if field == "report_id":
                cell = ws_bin.cell(row=r, column=c, value=row["report_id"])
                cell.font = cell_font(bold=True, size=9)
                cell.alignment = lft()
            elif field == "total_informative":
                cell = ws_bin.cell(row=r, column=c, value=total)
                cell.font = cell_font(bold=True, size=9)
                cell.fill = fill(COL_TOTAL)
                cell.alignment = ctr()
            else:
                val = row[field]
                cell = ws_bin.cell(row=r, column=c, value=val)
                cell.fill = fill(COL_PRESENT) if val == 1 else fill(COL_ABSENT)
                cell.font = cell_font(size=8)
                cell.alignment = ctr()
            cell.border = border

    # Two summary rows: N Present then % Present
    for sr_offset, (label, compute) in enumerate([
        ("N Present", lambda field, rows: sum(row[field] for row in rows)),
        ("% Present", lambda field, rows: round(sum(row[field] for row in rows) / n_reports * 100, 1)),
    ]):
        sr = last_data_row + 1 + sr_offset
        ws_bin.row_dimensions[sr].height = 15
        label_cell = ws_bin.cell(row=sr, column=1, value=label)
        label_cell.font = cell_font(bold=True, size=9)
        label_cell.fill = fill(COL_SUMROW)
        label_cell.alignment = lft()
        label_cell.border = border

        for c, field in enumerate(col_fields[1:], 2):
            if field == "total_informative":
                cell = ws_bin.cell(row=sr, column=c, value="")
            else:
                cell = ws_bin.cell(row=sr, column=c, value=compute(field, rows))
            cell.fill = fill(COL_SUMROW)
            cell.font = cell_font(bold=True, size=8)
            cell.alignment = ctr()
            cell.border = border

    ws_bin.freeze_panes = "B2"

    # ════════════════════════════════════════════════════════════════════════
    # SHEET 3 — Raw Values
    # None cells = all seeds returned None for that variable.
    # ════════════════════════════════════════════════════════════════════════
    ws_raw = wb.create_sheet("Raw Values")
    ws_raw.sheet_view.showGridLines = False

    raw_col_fields = ["report_id"] + active_vars

    # Header row — same style as Binary Matrix
    ws_raw.row_dimensions[1].height = 15
    for c, h in enumerate(raw_col_fields, 1):
        cell = ws_raw.cell(row=1, column=c, value=h)
        cell.fill = fill(HDR_DARK)
        cell.font = hdr_font(size=8)
        cell.alignment = Alignment(horizontal="center", vertical="center",
                                   wrap_text=False)
        cell.border = border
        if h == "report_id":
            ws_raw.column_dimensions[get_column_letter(c)].width = 22
        else:
            ws_raw.column_dimensions[get_column_letter(c)].width = 9

    # Data rows
    for r, raw_row in enumerate(raw_rows, 2):
        ws_raw.row_dimensions[r].height = 15
        for c, field in enumerate(raw_col_fields, 1):
            val = raw_row[field]
            cell = ws_raw.cell(row=r, column=c, value=val)
            cell.border = border
            if field == "report_id":
                cell.font = cell_font(bold=True, size=9)
                cell.alignment = lft()
            else:
                cell.font = cell_font(size=8)
                cell.alignment = ctr()
                # Colour: None = light red (no extraction), value = white
                if val is None:
                    cell.fill = fill("FADBD8")
                else:
                    cell.fill = fill(COL_ABSENT)  # white — let the value speak

    ws_raw.freeze_panes = "B2"

    wb.save(output_path)
    print(f"Excel written to: {output_path}")


def main():
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

    print(f"Mode:     {MODE}")
    print(f"Scanning: {INPUT_BASE}")
    reports = load_seed_files(INPUT_BASE)

    if not reports:
        print(f"No {'*.extracted.json' if MODE == 'primary' else '*.residual.json'} "
              f"files found in {INPUT_BASE}")
        return

    n_reports = len(reports)
    print(f"Found {n_reports} report(s)")
    print(f"Included variables: {len(_active_vars())}")
    if MODE == "primary":
        print(f"Excluded variables: {len(EXCLUDED_VARS)}")

    rows = []
    raw_rows = []
    for report_id in sorted(reports.keys()):
        seed_dicts = reports[report_id]
        n = len(seed_dicts)
        if n < N_SEEDS:
            print(f"  WARNING: {report_id} has only {n}/{N_SEEDS} seeds")
        row = compute_row(report_id, seed_dicts)
        raw_row = compute_raw_row(report_id, seed_dicts)
        rows.append(row)
        raw_rows.append(raw_row)
        informative_count = sum(v for k, v in row.items() if k != "report_id")
        print(f"  {report_id}: {informative_count}/{len(_active_vars())} vars informative")

    summary = build_variable_summary(rows, n_reports)

    output_path = OUTPUT_DIR / f"variable_matrix_{MODEL}_{MODE}.xlsx"
    write_excel(rows, raw_rows, summary, n_reports, output_path)

    # Quick console summary
    cats = {}
    for s in summary:
        cats.setdefault(s["Category"], 0)
        cats[s["Category"]] += 1
    print("\nVariable categories:")
    for cat, count in sorted(cats.items()):
        print(f"  {cat:<25} {count}")


if __name__ == "__main__":
    main()